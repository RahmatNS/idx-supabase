import os
from pathlib import Path

import openpyxl
import yfinance as yf
from supabase import create_client, Client

url = os.environ.get("SUPABASE_URL")
key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY") or os.environ.get("SUPABASE_KEY")
if not url or not key:
    raise RuntimeError("SUPABASE_URL dan SUPABASE_KEY/SUPABASE_SERVICE_ROLE_KEY wajib disetel")
supabase: Client = create_client(url, key)

STOCK_LIST_FILE = Path(__file__).with_name("Stock List.xlsx")


def get_tickers():
    if not STOCK_LIST_FILE.exists():
        raise FileNotFoundError(f"File daftar saham tidak ditemukan: {STOCK_LIST_FILE}")

    worksheet = openpyxl.load_workbook(STOCK_LIST_FILE, read_only=True, data_only=True).active
    headers = [cell.value for cell in next(worksheet.iter_rows())]
    try:
        code_column = headers.index("Code")
    except ValueError as error:
        raise RuntimeError("Kolom 'Code' tidak ditemukan di Stock List.xlsx") from error

    codes = {
        str(row[code_column].value).strip().upper()
        for row in worksheet.iter_rows(min_row=2)
        if row[code_column].value
    }
    if not codes:
        raise RuntimeError("Tidak ada kode saham di Stock List.xlsx")
    return ["^JKSE", *(f"{code}.JK" for code in sorted(codes))]

def main():
    tickers = get_tickers()
    print(f"Ditemukan {len(tickers) - 1} ticker dari Stock List.xlsx")
    prices = yf.download(tickers, period="1d", group_by="ticker", threads=True)
    payloads = []

    for ticker in tickers:
        try:
            history = prices[ticker].dropna(how="all")
            row = history.iloc[-1]
            payloads.append({
                "ticker": ticker,
                "date": history.index[-1].strftime("%Y-%m-%d"),
                "open": float(row["Open"]),
                "high": float(row["High"]),
                "low": float(row["Low"]),
                "close": float(row["Close"]),
                "volume": int(row["Volume"]),
            })
        except (KeyError, IndexError, TypeError, ValueError):
            print(f"Lewati: tidak ada harga untuk {ticker}")

    if not payloads:
        raise RuntimeError("Yahoo Finance tidak mengembalikan harga")
    response = supabase.table("ihsg_daily").upsert(
        payloads, on_conflict="ticker,date"
    ).execute()
    if not response.data:
        raise RuntimeError("Supabase tidak mengembalikan data setelah upsert")
    print(f"Disimpan: {len(response.data)} ticker")

if __name__ == "__main__":
    main()