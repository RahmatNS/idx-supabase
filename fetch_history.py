import os
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import pandas as pd
import yfinance as yf
from supabase import create_client, Client

url = os.environ.get("SUPABASE_URL")
key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY") or os.environ.get("SUPABASE_KEY")
if not url or not key:
    raise RuntimeError("SUPABASE_URL dan SUPABASE_KEY/SUPABASE_SERVICE_ROLE_KEY wajib disetel")

supabase: Client = create_client(url, key)
STOCK_LIST_FILE = Path(__file__).with_name("Stock List.xlsx")

# Target maksimal backfill (5 tahun ke belakang)
MAX_YEARS = 5
CHUNK_DAYS = 30
BATCH_UPSERT_SIZE = 1000


def get_tickers():
    if not STOCK_LIST_FILE.exists():
        raise FileNotFoundError(f"File daftar saham tidak ditemukan: {STOCK_LIST_FILE}")

    worksheet = openpyxl.load_workbook(STOCK_LIST_FILE, read_only=True, data_only=True).active
    headers = [cell.value for cell in next(worksheet.iter_rows())]
    try:
        code_column = headers.index("Code")
    except ValueError as error:
        raise RuntimeError("Kolom 'Code' tidak ditemukan di Stock List.xlsx") from error

    codes = {
        str(row[code_column].value).strip().upper()
        for row in worksheet.iter_rows(min_row=2)
        if row[code_column].value
    }
    if not codes:
        raise RuntimeError("Tidak ada kode saham di Stock List.xlsx")
    return ["^JKSE", *(f"{code}.JK" for code in sorted(codes))]


def get_earliest_date_in_db() -> date | None:
    try:
        res = (
            supabase.table("ihsg_daily")
            .select("date")
            .order("date", desc=False)
            .limit(1)
            .execute()
        )
        if res.data and len(res.data) > 0 and res.data[0].get("date"):
            return datetime.strptime(res.data[0]["date"], "%Y-%m-%d").date()
    except Exception as e:
        print(f"Peringatan: Gagal mengecek tanggal di database ({e}), menggunakan hari ini.")
    return None


def main():
    today = datetime.now().date()
    cutoff_date = today - timedelta(days=MAX_YEARS * 365)

    earliest_date = get_earliest_date_in_db()
    if earliest_date is None:
        # Jika belum ada data sama sekali di DB, gunakan hari ini sebagai titik akhir
        end_date = today
    else:
        end_date = earliest_date

    # Jika sudah mencapai batas 5 tahun ke belakang, keluar tanpa fetch lagi
    if end_date <= cutoff_date:
        sys.exit(0)

    start_date = max(end_date - timedelta(days=CHUNK_DAYS), cutoff_date)
    if start_date >= end_date:
        sys.exit(0)

    tickers = get_tickers()
    print("Memulai backfill mundur...")
    print(f"Titik awal di database: {earliest_date}")
    print(f"Jendela fetch: {start_date} s/d {end_date} (semua {len(tickers)} ticker)")

    # yfinance end date bersifat eksklusif, gunakan string YYYY-MM-DD
    # Download semua ticker sekaligus untuk jendela waktu ini
    df = yf.download(
        tickers=tickers,
        start=start_date.strftime("%Y-%m-%d"),
        end=end_date.strftime("%Y-%m-%d"),
        group_by="ticker",
        threads=True,
        auto_adjust=False,
    )

    if df.empty:
        print("Tidak ada data dikembalikan oleh Yahoo Finance untuk jendela waktu ini.")
        return

    payloads = []
    is_multi_ticker = isinstance(df.columns, pd.MultiIndex)

    for ticker in tickers:
        try:
            if is_multi_ticker:
                if ticker not in df.columns.levels[0]:
                    continue
                sub_df = df[ticker].dropna(how="all")
            else:
                sub_df = df.dropna(how="all")

            for timestamp, row in sub_df.iterrows():
                # Pastikan nilai numerik valid dan tidak NaN
                if pd.isna(row.get("Close")) or pd.isna(row.get("Volume")):
                    continue

                date_str = timestamp.strftime("%Y-%m-%d")
                payloads.append({
                    "ticker": ticker,
                    "date": date_str,
                    "open": float(row["Open"]),
                    "high": float(row["High"]),
                    "low": float(row["Low"]),
                    "close": float(row["Close"]),
                    "volume": int(row["Volume"]),
                })
        except Exception:
            continue

    if not payloads:
        print("Tidak ada baris data valid untuk disimpan.")
        return

    total_rows = len(payloads)
    print(f"Total baris valid yang ditemukan: {total_rows}. Mengirim ke Supabase per {BATCH_UPSERT_SIZE} baris...")

    # Batch upsert ke Supabase
    total_inserted = 0
    for i in range(0, total_rows, BATCH_UPSERT_SIZE):
        batch = payloads[i:i + BATCH_UPSERT_SIZE]
        try:
            supabase.table("ihsg_daily").upsert(batch, on_conflict="ticker,date").execute()
            total_inserted += len(batch)
            pct = (total_inserted / total_rows) * 100
            print(f"Progress: {total_inserted}/{total_rows} ({pct:.1f}%) baris tersimpan")
        except Exception as e:
            print(f"Error saat upsert batch {i} - {i + len(batch)}: {e}")

    print(f"Selesai! Berhasil menyimpan data rentang {start_date} s/d {end_date}.")


if __name__ == "__main__":
    main()
